from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document as DocxDocument

def extract_text(filename: str, data: bytes) -> str:
    ext=Path(filename).suffix.lower()
    if ext=='.pdf':
        import io
        reader=PdfReader(io.BytesIO(data)); return '\n\n'.join((p.extract_text(extraction_mode='layout') or '') for p in reader.pages)
    if ext=='.docx':
        import io
        d=DocxDocument(io.BytesIO(data)); parts=[p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            parts.append('\n'.join(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows))
        return '\n\n'.join(parts)
    if ext in {'.txt','.md','.csv','.json','.xml','.html'}:
        return data.decode('utf-8', errors='ignore')
    raise ValueError(f'Unsupported document type: {ext}')

def chunk_text(text: str, size=1000, overlap=150):
    text='\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines())
    text=re.sub(r'\n{3,}', '\n\n', text).strip()
    chunks=[]; start=0
    while start < len(text):
        end=min(len(text), start+size); chunks.append(text[start:end])
        if end==len(text): break
        start=max(0,end-overlap)
    return chunks

def extract_salary_scale_markdown(data: bytes) -> str | None:
    import io

    def numeric_values(text: str) -> list[str]:
        matches=re.findall(r'(?<!\d)\d{1,3}(?:[ \u00a0]\d{3})+|(?<!\d)\d+(?!\d)', text)
        return [' '.join(value.split()) for value in matches]

    reader=PdfReader(io.BytesIO(data))
    for page in reader.pages:
        if 'salary scale' not in (page.extract_text() or '').lower():
            continue
        words=[]
        page.extract_text(visitor_text=lambda text, cm, tm, font, size: words.append((tm[4],tm[5],text)))
        rows=[]
        for x,y,text in words:
            value=text.strip()
            if value in {str(number) for number in range(4,13)} and x < 140:
                rows.append((value,y))
        if not rows:
            continue
        output=['| Echelon | Minimum | Médian | Maximum |','| --- | --- | --- | --- |']
        for echelon,y in rows:
            cells=[]
            for x,cell_y,text in words:
                if x >= 145 and abs(cell_y-y) <= 5 and text.strip():
                    cells.extend((x,value) for value in numeric_values(text))
            values=[value for _,value in sorted(cells)]
            if values:
                middle=values[(len(values)-1)//2]
                output.append(f'| {echelon} | {values[0]} | {middle} | {values[-1]} |')
        return '\n'.join(output) if len(output)>2 else None
    return None
